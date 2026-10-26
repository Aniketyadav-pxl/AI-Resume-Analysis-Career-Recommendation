"""M1 - Resume text extraction from PDF, DOCX and TXT.

All failures are raised as ExtractionError with a user-friendly message, so the UI
can show it directly without crashing.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
MAX_PAGES = 10
MAX_CHARS = 100_000
MIN_CHARS = 50
SUPPORTED = (".pdf", ".docx", ".txt")


class ExtractionError(Exception):
    """Raised when a resume file cannot be turned into usable text."""


@dataclass
class ExtractionResult:
    text: str
    file_type: str
    filename: str = ""
    n_pages: int = 1
    warnings: list = field(default_factory=list)


def detect_type(filename: str, data: bytes) -> str:
    """Decide the real type from magic bytes first, falling back to the extension."""
    ext = Path(filename or "").suffix.lower()
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":  # zip container -> DOCX (we verify when opening)
        return "docx"
    if ext == ".doc":
        raise ExtractionError("Old .doc files are not supported. Please save the resume as .docx or PDF.")
    if ext in SUPPORTED:
        if ext in (".pdf", ".docx"):
            raise ExtractionError(f"This file is named '{ext}' but its content is not a valid {ext[1:].upper()} file.")
        return "txt"
    raise ExtractionError(f"Unsupported file type '{ext or 'unknown'}'. Upload a PDF, DOCX or TXT file.")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_pdf(data: bytes, warnings: list) -> tuple:
    import pymupdf  # imported lazily so the rest of the app works without it

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError("Could not open this PDF. The file may be corrupted.") from exc
    try:
        if doc.needs_pass:
            raise ExtractionError("This PDF is password-protected. Remove the password and upload again.")
        n_pages = doc.page_count
        if n_pages > MAX_PAGES:
            warnings.append(f"PDF has {n_pages} pages; only the first {MAX_PAGES} were read.")
        pages = [doc[i].get_text("text", sort=True) for i in range(min(n_pages, MAX_PAGES))]
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Could not read text from this PDF.") from exc
    finally:
        doc.close()
    return "\n".join(pages), n_pages


def _from_docx(data: bytes, warnings: list) -> tuple:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("Could not open this DOCX file. It may be corrupted or not a Word document.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts), 1


def _from_txt(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def extract_text(data: bytes, filename: str = "") -> ExtractionResult:
    """Extract clean text from raw file bytes."""
    if not data:
        raise ExtractionError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ExtractionError(f"File is too large ({len(data) / 1e6:.1f} MB). Maximum size is {MAX_BYTES // 1_000_000} MB.")

    file_type = detect_type(filename, data)
    warnings: list = []
    n_pages = 1
    if file_type == "pdf":
        text, n_pages = _from_pdf(data, warnings)
    elif file_type == "docx":
        text, n_pages = _from_docx(data, warnings)
    else:
        text = _from_txt(data)

    text = clean_text(text)
    if len(text) < MIN_CHARS:
        hint = " It may be a scanned image; OCR is not supported." if file_type == "pdf" else ""
        raise ExtractionError(f"Very little text could be extracted ({len(text)} characters).{hint}")
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS]
        warnings.append("Resume text was truncated to 100,000 characters.")
    return ExtractionResult(text=text, file_type=file_type, filename=filename, n_pages=n_pages, warnings=warnings)


def extract_from_path(path) -> ExtractionResult:
    path = Path(path)
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ExtractionError(f"Could not read file: {path.name}") from exc
    return extract_text(data, path.name)
