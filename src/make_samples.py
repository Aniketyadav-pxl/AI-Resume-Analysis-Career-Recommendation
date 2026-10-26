"""Create synthetic sample resumes (TXT, DOCX, PDF) in sample_resumes/.

These use a different random seed from the training data, so they are unseen by the model.
Run:  python -m src.make_samples
"""
import random

import config
from src.data_generator import build_resume

PLAN = [  # (role, extension)
    ("Data Scientist", "txt"), ("Frontend Developer", "pdf"), ("DevOps Engineer", "docx"),
    ("Cybersecurity Analyst", "pdf"), ("Mobile App Developer", "docx"), ("Data Analyst", "txt"),
    ("QA / Test Engineer", "txt"), ("Backend Developer", "pdf"), ("Machine Learning Engineer", "docx"),
    ("Full Stack Developer", "pdf"),
]
HEADERS = {"SUMMARY", "SKILLS", "EXPERIENCE", "PROJECTS", "EDUCATION"}


def _slug(role: str) -> str:
    return role.lower().replace(" / ", "_").replace(" ", "_")


def write_pdf(text: str, path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page, y = doc.new_page(), 60
    for line in text.split("\n"):
        if y > 780:
            page, y = doc.new_page(), 60
        bold = line in HEADERS or line.startswith("CANDIDATE")
        page.insert_text((50, y), line, fontsize=12 if bold else 10, fontname="hebo" if bold else "helv")
        y += 18 if bold else 14
    doc.save(str(path))
    doc.close()


def write_docx(text: str, path) -> None:
    import docx

    d = docx.Document()
    for line in text.split("\n"):
        if line in HEADERS:
            d.add_heading(line.title(), level=2)
        else:
            d.add_paragraph(line)
    d.save(str(path))


def main() -> None:
    config.SAMPLE_DIR.mkdir(exist_ok=True)
    rng = random.Random(2025)
    for i, (role, ext) in enumerate(PLAN):
        text = build_resume(9001 + i, role, rng)["resume_text"]
        path = config.SAMPLE_DIR / f"sample_{_slug(role)}.{ext}"
        if ext == "txt":
            path.write_text(text, encoding="utf-8")
        elif ext == "docx":
            write_docx(text, path)
        else:
            write_pdf(text, path)
        print("wrote", path.name)


if __name__ == "__main__":
    main()
