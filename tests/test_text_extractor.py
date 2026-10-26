import io

import docx
import pytest

import config
from src.text_extractor import ExtractionError, extract_from_path, extract_text

SAMPLES = sorted(config.SAMPLE_DIR.glob("sample_*.*"))
LONG = "Experienced developer with Python and SQL skills. " * 5


def test_all_generated_samples_extract():
    assert len(SAMPLES) >= 10
    types = set()
    for p in SAMPLES:
        r = extract_from_path(p)
        types.add(r.file_type)
        assert "SKILLS" in r.text.upper() and len(r.text) > 300, p.name
    assert types == {"pdf", "docx", "txt"}


def test_txt_encodings():
    assert "caf" in extract_text(("café " + LONG).encode("utf-8"), "a.txt").text
    assert "café" in extract_text(("café " + LONG).encode("cp1252"), "a.txt").text
    assert "Python" in extract_text(LONG.encode("utf-16"), "a.txt").text


def test_docx_tables_are_read():
    d = docx.Document()
    d.add_paragraph(LONG)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Skill", "Docker"
    buf = io.BytesIO()
    d.save(buf)
    assert "Skill | Docker" in extract_text(buf.getvalue(), "r.docx").text


@pytest.mark.parametrize("data,name,msg", [
    (b"", "a.txt", "empty"),
    (b"x" * (6 * 1024 * 1024), "a.txt", "too large"),
    (b"hello", "a.txt", "little text"),
    (b"data", "a.exe", "Unsupported"),
    (b"not really a pdf" * 10, "a.pdf", "not a valid PDF"),
    (b"%PDF-1.4 garbage garbage", "a.pdf", "PDF"),
    (b"PK\x03\x04 not a docx", "a.docx", "DOCX"),
    (b"anything", "old.doc", ".docx"),
])
def test_errors_are_friendly(data, name, msg):
    with pytest.raises(ExtractionError) as e:
        extract_text(data, name)
    assert msg.lower() in str(e.value).lower()


def test_wrong_extension_uses_magic_bytes():
    pdf = next(p for p in SAMPLES if p.suffix == ".pdf").read_bytes()
    assert extract_text(pdf, "renamed.txt").file_type == "pdf"


def test_empty_pdf_reports_no_text():
    import pymupdf
    doc = pymupdf.open()
    doc.new_page()
    with pytest.raises(ExtractionError) as e:
        extract_text(doc.tobytes(), "scan.pdf")
    assert "OCR" in str(e.value)


def test_missing_file_path():
    with pytest.raises(ExtractionError):
        extract_from_path(config.SAMPLE_DIR / "nope.pdf")
